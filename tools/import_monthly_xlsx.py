# -*- coding: utf-8 -*-
"""월말 평가 엑셀(월별 시트 2026-01 ...) -> data/history/{kr,jp}/YYYY-MM.json 스냅샷 변환

  python tools/import_monthly_xlsx.py 한국_여자아이돌_월말평가_2026-01_08.xlsx 일본_여자아이돌_월말평가_2026-01_08.xlsx
  python tools/import_monthly_xlsx.py --dry-run <xlsx...>     # 검증만 하고 쓰지 않는다

원칙 (build_history.py 와 동일)
- 시트 이름이 YYYY-MM 인 시트만 읽는다. 순위·총점·등급은 엑셀 값을 그대로 저장한다 (RANK.EQ 공동순위 유지).
- 그룹 식별은 DB 의 id (이름은 NFKC 정규화로 매칭). 이름이 DB에 없으면 중단한다.
- 지표 6개 합계가 총점과 다르면 중단한다. 그 달에 없는 그룹(데뷔 전 등)은 그냥 빠진다 (population_count 에 반영).
- 이미 있는 달은 덮어쓰지 않는다 (정정은 build_history.py --correct).
"""
import argparse
import calendar
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_history as bh  # noqa: E402

KR_HEAD = "국내음원"
JP_HEAD = "공연·현장"


def norm(s):
    return unicodedata.normalize("NFKC", str(s)).strip()


def head_of(ws):
    for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        if r and r[0] == "순위" and r[1] == "그룹":
            return r
    return None


def read_sheet(ws, metrics, name_to_id, label):
    head = head_of(ws)
    if not head:
        sys.exit("%s: 헤더(순위/그룹)를 찾지 못했습니다" % label)
    first = norm(str(head[2]).split("\n")[0])
    if first != metrics[0][0]:
        sys.exit("%s: 첫 지표 '%s' 가 예상 '%s' 와 다릅니다 (한국/일본 파일이 바뀐 것 같습니다)" % (label, first, metrics[0][0]))
    out, seen = [], set()
    started = False
    for r in ws.iter_rows(values_only=True):
        if not started:
            started = r is head or (r and r[0] == "순위" and r[1] == "그룹")
            continue
        if r[0] is None or r[1] is None:
            continue
        nm = norm(r[1])
        if nm not in name_to_id:
            sys.exit("%s: DB에 없는 그룹 '%s'" % (label, r[1]))
        gid = name_to_id[nm]
        if gid in seen:
            sys.exit("%s: 그룹 중복 '%s'" % (label, r[1]))
        seen.add(gid)
        vals = [r[i] for i in range(2, 8)]
        if any(not isinstance(v, (int, float)) for v in vals) or not isinstance(r[8], (int, float)):
            sys.exit("%s %s: 숫자가 아닌 값이 있습니다 %s" % (label, r[1], vals))
        if abs(sum(vals) - r[8]) > 0.01:
            sys.exit("%s %s: 지표 합 %s != 총점 %s" % (label, r[1], sum(vals), r[8]))
        out.append({"id": gid, "group": r[1], "rank": int(r[0]), "tier": r[9], "score": r[8],
                    "metrics": {key: vals[i] for i, (_, key) in enumerate(metrics)}})
    return out


def check_ranks(groups, label):
    order = sorted(groups, key=lambda g: -g["score"])
    exp, prev, cur = {}, None, 0
    for i, g in enumerate(order):
        if g["score"] != prev:
            cur, prev = i + 1, g["score"]
        exp[g["id"]] = cur
    bad = [g["group"] for g in groups if g["rank"] != exp[g["id"]]]
    if bad:
        print("  경고 %s: 엑셀 순위가 competition rank(1,2,2,4)와 다른 그룹 %d개 (엑셀 값을 그대로 저장): %s" % (label, len(bad), ", ".join(bad[:5])))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="+")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    files = []
    for p in a.xlsx:
        files += glob.glob(p) or [p]

    for path in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        periods = [n for n in wb.sheetnames if re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", n)]
        if not periods:
            sys.exit("%s: YYYY-MM 시트가 없습니다" % path)
        first = head_of(wb[periods[0]])
        country = "KR" if first and norm(str(first[2]).split("\n")[0]) == KR_HEAD else "JP" if first and norm(str(first[2]).split("\n")[0]) == JP_HEAD else None
        if not country:
            sys.exit("%s: 한국/일본 형식을 알아볼 수 없습니다" % path)
        metrics = bh.KR_METRICS if country == "KR" else bh.JP_METRICS
        rows = bh.load_rows(os.path.join(bh.ROOT, "data", "%s_db.json" % country.lower()), country)
        name_to_id = {norm(r["그룹"]): r["id"] for r in rows}
        for period in periods:
            label = "%s %s" % (country, period)
            groups = read_sheet(wb[period], metrics, name_to_id, label)
            check_ranks(groups, label)
            groups.sort(key=lambda g: (g["rank"], g["id"]))
            y, m = int(period[:4]), int(period[5:])
            published = "%s-%02d" % (period, calendar.monthrange(y, m)[1])  # 월말 기준일
            meta = {"country": country, "period": period, "published_at": published, "revision": 1,
                    "population_count": len(groups), "normalization_version": "v2",
                    "percentile_scope": "same-country, same-period relative position",
                    "source": "월말 평가 엑셀(후향적 상대평가, 해당 월 말일 23:59 KST까지 공개된 자료만 반영)"}
            out = os.path.join(bh.HIST, country.lower(), period + ".json")
            if os.path.exists(out):
                print("건너뜀(이미 있음): %s" % os.path.relpath(out, bh.ROOT))
                continue
            print("%s: %d teams%s" % (label, len(groups), " (dry-run)" if a.dry_run else " -> " + os.path.relpath(out, bh.ROOT)))
            if not a.dry_run:
                bh.write_json(out, {"meta": meta, "groups": groups})
                bh.update_index(country, period)


if __name__ == "__main__":
    main()
